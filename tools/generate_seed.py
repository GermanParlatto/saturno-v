#!/usr/bin/env python3
"""
generate_seed.py — Genera el SQL de seed de Waku Code desde los ficheros fuente.

Los CSV y el XLSX NO están versionados (viven fuera del repo). Este script los
procesa una vez y emite dos migraciones SQL que sí se commitean, de forma que el
seed sea reproducible en cualquier entorno sin arrastrar los ficheros de origen.

Reglas de negocio (verificadas contra los datos):
  * tb_nodes.csv es la fuente de verdad de QUÉ nodos existen: los 70, todos.
  * assets.xlsx solo AÑADE stage/mission/objective/metadata, y solo desde filas
    cuyo `status` NO sea 'Descartado'. Casan 42 de los 70.
  * Los 28 sub-nodos restantes (E1-01-01, etc.) derivan stage/mission del prefijo
    del id. OJO: las filas R0-* tienen stage/mission VACÍOS en la hoja (no 0), y
    no siguen el patrón E<n>, así que se quedan en NULL.
  * No hay conflictos de file_url entre CSV y XLSX: el enriquecimiento es aditivo.

Uso:
  make seed-sql
  python tools/generate_seed.py --nodes ../tb_nodes.csv \
      --flow ../tb_flow_sequence.csv --assets ../assets.xlsx
"""

import argparse
import csv
import json
import pathlib
import sys

# OJO: openpyxl NO se importa aqui, sino dentro de leer_assets(). Vive en el
# grupo `tools`, que el CI no instala (`uv sync --locked --group dev`), asi que
# un import a nivel de modulo romperia la recoleccion de tests. Con el import
# diferido, las funciones puras (lit, etapa_desde_id, normaliza_metadata) se
# pueden importar y testear sin la dependencia.

OUT_DIR = pathlib.Path(__file__).resolve().parents[1] / "src" / "migrations"

CABECERA = (
    "-- GENERADO por tools/generate_seed.py - no editar a mano.\n"
    "-- Fuente: {fuentes}\n"
    "-- Regenerar con: make seed-sql\n"
)

# stage == mission == el dígito tras la 'E'. EF es la etapa final (5).
PREFIJO_ETAPA = {"E0": 0, "E1": 1, "E2": 2, "E3": 3, "E4": 4, "EF": 5}


def lit(value) -> str:
    """Convierte un valor Python a un literal SQL.

    El escapado NO se improvisa en el sitio de uso: las descripciones son prosa
    en español llena de apóstrofes y con saltos de línea embebidos. Un solo
    escape mal puesto rompe la migración entera.
    """
    if value is None or value == "":
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def etapa_desde_id(node_id: str) -> int | None:
    """stage/mission derivados del prefijo, para los nodos sin fila en assets."""
    return PREFIJO_ETAPA.get(node_id[:2].upper())


def leer_nodos(ruta: pathlib.Path) -> list[dict]:
    # utf-8-sig: los CSV vienen con BOM. Con utf-8 a secas, la primera columna se
    # llamaría '﻿id' y todos los lookups por 'id' fallarían.
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def leer_assets(ruta: pathlib.Path) -> dict[str, dict]:
    """Filas de la hoja 'assets' indexadas por id, excluyendo las descartadas."""
    # Import diferido: ver la nota de la cabecera.
    from openpyxl import load_workbook

    libro = load_workbook(ruta, data_only=True, read_only=True)
    hoja = libro["assets"]
    filas = hoja.iter_rows(values_only=True)
    cabecera = [str(c).strip() if c is not None else "" for c in next(filas)]

    resultado: dict[str, dict] = {}
    for fila in filas:
        registro = dict(zip(cabecera, fila, strict=False))
        node_id = (registro.get("id") or "").strip()
        if not node_id:
            continue
        # Los assets descartados no se cargan ni enriquecen nada.
        if (registro.get("status") or "").strip() == "Descartado":
            continue
        resultado[node_id] = registro
    libro.close()
    return resultado


def normaliza_metadata(bruto) -> str:
    """`metadata` viene como texto JSON en la hoja. Se valida aquí para no
    generar SQL que falle al hacer el CAST a jsonb."""
    if not bruto:
        return "{}"
    if isinstance(bruto, dict):
        return json.dumps(bruto, ensure_ascii=False)
    try:
        return json.dumps(json.loads(str(bruto)), ensure_ascii=False)
    except (ValueError, TypeError):
        print(f"  aviso: metadata no es JSON válido, se ignora: {bruto!r}", file=sys.stderr)
        return "{}"


def construye_filas_nodos(nodos: list[dict], assets: dict[str, dict]) -> list[tuple]:
    filas = []
    enriquecidos = 0
    for nodo in nodos:
        node_id = nodo["id"].strip()
        extra = assets.get(node_id)

        if extra is not None:
            enriquecidos += 1
            stage = extra.get("stage")
            mission = extra.get("mission")
            objective = (extra.get("objective") or "").strip() or None
            metadata = normaliza_metadata(extra.get("metadata"))
            # Vacío en la hoja (filas R0-*) significa "sin etapa", no 0.
            stage = int(stage) if stage not in (None, "") else None
            mission = int(mission) if mission not in (None, "") else None
        else:
            stage = mission = etapa_desde_id(node_id)
            objective = None
            metadata = "{}"

        filas.append(
            (
                node_id,
                nodo["type"].strip(),
                nodo["function"].strip(),
                nodo["pauses"].strip().upper() == "TRUE",
                (nodo["description"] or "").strip() or None,
                (nodo["Output"] or "").strip() or None,
                (nodo["Logic"] or "").strip() or None,
                (nodo["file_url"] or "").strip() or None,
                stage,
                mission,
                objective,
                metadata,
            )
        )
    print(f"  nodos: {len(filas)} ({enriquecidos} enriquecidos desde assets.xlsx)")
    return filas


def sql_nodos(filas: list[tuple], fuentes: str) -> str:
    valores = ",\n".join(
        "    ({})".format(
            ", ".join(
                [
                    lit(f[0]),
                    lit(f[1]),
                    lit(f[2]),
                    lit(f[3]),
                    lit(f[4]),
                    lit(f[5]),
                    lit(f[6]),
                    lit(f[7]),
                    lit(f[8]),
                    lit(f[9]),
                    lit(f[10]),
                    f"{lit(f[11])}::jsonb",
                ]
            )
        )
        for f in filas
    )
    # DO UPDATE, no DO NOTHING: regenerar el fichero tras cambiar la hoja tiene
    # que propagar el cambio, no quedarse en un no-op silencioso.
    return (
        CABECERA.format(fuentes=fuentes)
        + "\nINSERT INTO nodes (id, type, function, pauses, description, output, logic,\n"
        "                   file_url, stage, mission, objective, metadata) VALUES\n"
        + valores
        + "\nON CONFLICT (id) DO UPDATE SET\n"
        "    type        = EXCLUDED.type,\n"
        "    function    = EXCLUDED.function,\n"
        "    pauses      = EXCLUDED.pauses,\n"
        "    description = EXCLUDED.description,\n"
        "    output      = EXCLUDED.output,\n"
        "    logic       = EXCLUDED.logic,\n"
        "    file_url    = EXCLUDED.file_url,\n"
        "    stage       = EXCLUDED.stage,\n"
        "    mission     = EXCLUDED.mission,\n"
        "    objective   = EXCLUDED.objective,\n"
        "    metadata    = EXCLUDED.metadata,\n"
        "    updated_at  = now();\n"
    )


def sql_flujo(flujo: list[dict], ids_validos: set[str], fuentes: str) -> str:
    filas = []
    for fila in flujo:
        node_id = fila["node_id"].strip()
        # La FK lo impediría en la BD, pero fallar aquí da un error legible en
        # vez de una violación de clave ajena a mitad de la migración.
        if node_id not in ids_validos:
            raise SystemExit(f"ERROR: flow_sequence referencia un nodo inexistente: {node_id}")
        filas.append((int(fila["id"]), int(fila["order"]), node_id))

    valores = ",\n".join(f"    ({lit(a)}, {lit(b)}, {lit(c)})" for a, b, c in sorted(filas))
    print(f"  flow_sequence: {len(filas)} posiciones")
    return (
        CABECERA.format(fuentes=fuentes)
        + "\n-- La columna `order` del CSV es palabra reservada en SQL: aqui es step_order.\n"
        "INSERT INTO flow_sequence (id, step_order, node_id) VALUES\n"
        + valores
        + "\nON CONFLICT (id) DO UPDATE SET\n"
        "    step_order = EXCLUDED.step_order,\n"
        "    node_id    = EXCLUDED.node_id;\n"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--nodes", required=True, type=pathlib.Path)
    parser.add_argument("--flow", required=True, type=pathlib.Path)
    parser.add_argument("--assets", required=True, type=pathlib.Path)
    parser.add_argument("--out", type=pathlib.Path, default=OUT_DIR)
    args = parser.parse_args()

    for ruta in (args.nodes, args.flow, args.assets):
        if not ruta.exists():
            raise SystemExit(f"ERROR: no existe {ruta}")

    fuentes = f"{args.nodes.name} + {args.flow.name} + {args.assets.name} (status != 'Descartado')"

    print("Generando seed SQL...")
    nodos = leer_nodos(args.nodes)
    assets = leer_assets(args.assets)
    flujo = leer_nodos(args.flow)

    filas_nodos = construye_filas_nodos(nodos, assets)
    ids = {f[0] for f in filas_nodos}
    if len(ids) != len(filas_nodos):
        raise SystemExit("ERROR: hay ids de nodo duplicados en el CSV")

    args.out.mkdir(parents=True, exist_ok=True)
    destino_nodos = args.out / "003_seed_nodes.sql"
    destino_flujo = args.out / "004_seed_flow_sequence.sql"

    destino_nodos.write_text(sql_nodos(filas_nodos, fuentes), encoding="utf-8")
    destino_flujo.write_text(sql_flujo(flujo, ids, fuentes), encoding="utf-8")

    print(f"  escrito {destino_nodos}")
    print(f"  escrito {destino_flujo}")


if __name__ == "__main__":
    main()
